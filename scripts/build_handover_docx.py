from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
TABLE_FILL = "E8EEF5"
BORDER = "C9D4E2"


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, *, size: float, bold: bool = False, color: str = "000000") -> None:
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
        "Heading 4": (11, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        set_style_font(style, size=size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, size=11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Code Inline" not in [style.name for style in doc.styles]:
        code = doc.styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
        code.font.name = "Consolas"
        code._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
        code._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        code._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        code.font.size = Pt(9.5)
        code.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    configure_headers(section)


def configure_headers(section) -> None:
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("园区低碳规划与绿电直连优化平台  |  项目交接与演示手册")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    label = paragraph.add_run("v0.7.0  ·  ")
    set_run_font(label, size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_element = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    properties.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    properties.append(color)
    run_element.append(properties)
    text = OxmlElement("w:t")
    text.text = "1"
    run_element.append(text)
    field.append(run_element)
    paragraph._p.append(field)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(108)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("PROJECT HANDOVER · RELEASE 0.7.0")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("园区低碳规划与\n绿电直连优化平台")
    set_run_font(run, size=29, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("项目交接手册与逐页面演示讲稿")
    set_run_font(run, size=15, color=DARK_BLUE)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(84)
    run = summary.add_run("可部署  ·  可交付  ·  可维护  ·  可扩展")
    set_run_font(run, size=11, bold=True, color=BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    run = meta.add_run(f"版本日期：{date.today().isoformat()}")
    set_run_font(run, size=10, color=MUTED)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("交付分支：upgrade/v0.7.0-productization")
    set_run_font(run, size=10, color=MUTED)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("阅读导航", level=1)
    paragraph = doc.add_paragraph(
        "本文件由两部分组成：第一部分用于工程交接与维护，第二部分可直接作为现场逐页面讲稿。"
    )
    paragraph.paragraph_format.space_after = Pt(10)
    sections = [
        "产品定位、交付物与系统架构",
        "快速调度与V17两套逐页面说明",
        "优化、预测、比较、任务与备份算法",
        "数据字典、API、部署、测试与维护",
        "15分钟主线演示、全页面扩展讲法",
        "现场问答、故障应急与演示后动作",
    ]
    for item in sections:
        doc.add_paragraph(item, style="List Bullet")
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run("使用建议：")
    set_run_font(run, bold=True, color=DARK_BLUE)
    add_inline_runs(note, "演示人员先阅读第二部分；工程师按第一部分完成环境、模型、数据与运维交接。")
    doc.add_page_break()


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_inline_runs(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.style = "Code Inline"
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run)


def text_weight(value: str) -> int:
    return sum(2 if ord(char) > 127 else 1 for char in value)


def calculate_widths(rows: list[list[str]]) -> list[int]:
    column_count = len(rows[0])
    weights = []
    for index in range(column_count):
        maximum = max(text_weight(row[index]) for row in rows)
        weights.append(max(6, min(maximum, 32)))
    minimum = 850 if column_count >= 4 else 1050
    remaining = CONTENT_WIDTH_DXA - minimum * column_count
    total_weight = sum(weights)
    widths = [minimum + round(remaining * weight / total_weight) for weight in weights]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_cell_margins(table) -> None:
    table_properties = table._tbl.tblPr
    margins = table_properties.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        table_properties.append(margins)
    for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.first_child_found_in("w:tcW")
            cell_width.set(qn("w:w"), str(widths[index]))
            cell_width.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(table)


def set_cell_border(cell, color: str = BORDER) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.first_child_found_in("w:shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=column_count)
    table.style = "Table Grid"
    widths = calculate_widths(normalized)
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for row_index, (word_row, values) in enumerate(zip(table.rows, normalized)):
        prevent_row_split(word_row)
        for cell, value in zip(word_row.cells, values):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline_runs(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, size=9.2, bold=(row_index == 0), color=(NAVY if row_index == 0 else "000000"))
            set_cell_border(cell)
            if row_index == 0:
                shade_cell(cell, TABLE_FILL)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value.replace(" ", "")) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def create_decimal_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    number_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    number_id = max(number_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level.append(number_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    alignment = OxmlElement("w:lvlJc")
    alignment.set(qn("w:val"), "left")
    level.append(alignment)
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    paragraph_properties.append(indent)
    level.append(paragraph_properties)
    abstract.append(level)

    first_number = numbering.find(qn("w:num"))
    if first_number is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_number), abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return number_id


def apply_numbering(paragraph, number_id: int) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    number_properties = paragraph_properties.get_or_add_numPr()
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(number_id))
    number_properties.append(level)
    number_properties.append(number)


def add_markdown(doc: Document, path: Path, *, skip_title: bool = True) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    skipped = False
    paragraph_lines: list[str] = []
    current_numbering_id: int | None = None

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines)
        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, text)
        paragraph_lines.clear()

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            current_numbering_id = None
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            flush_paragraph()
            current_numbering_id = None
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            current_numbering_id = None
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1 and skip_title and not skipped:
                skipped = True
            else:
                doc.add_heading(title, level=min(level, 4))
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            flush_paragraph()
            if current_numbering_id is None:
                current_numbering_id = create_decimal_numbering(doc)
            paragraph = doc.add_paragraph(style="List Number")
            apply_numbering(paragraph, current_numbering_id)
            add_inline_runs(paragraph, numbered.group(1))
            index += 1
            continue
        current_numbering_id = None
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet.group(1))
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def build(output: Path) -> Path:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    part = doc.add_paragraph()
    part.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part.paragraph_format.space_before = Pt(16)
    part.paragraph_format.space_after = Pt(18)
    run = part.add_run("第一部分  ·  项目交接手册")
    set_run_font(run, size=19, bold=True, color=NAVY)
    add_markdown(doc, PROJECT_ROOT / "docs" / "project_handover_manual.md")

    part = doc.add_paragraph()
    part.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part.paragraph_format.page_break_before = True
    part.paragraph_format.space_before = Pt(24)
    part.paragraph_format.space_after = Pt(18)
    run = part.add_run("第二部分  ·  逐页面演示讲稿")
    set_run_font(run, size=19, bold=True, color=NAVY)
    add_markdown(doc, PROJECT_ROOT / "docs" / "demo_script.md")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "园区低碳规划与绿电直连优化平台 v0.7.0 项目交接与演示手册"
    doc.core_properties.subject = "项目交接、算法实现、部署维护与逐页面演示"
    doc.core_properties.author = "项目交付组"
    doc.core_properties.keywords = "园区低碳, 绿电直连, 综合能源, 项目交接, 演示讲稿"
    doc.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="生成v0.7.0项目交接与演示Word手册")
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "deliverables" / "园区低碳优化平台_v0.7.0_项目交接与演示手册.docx",
    )
    args = parser.parse_args()
    print(build(args.output.resolve()))


if __name__ == "__main__":
    main()
