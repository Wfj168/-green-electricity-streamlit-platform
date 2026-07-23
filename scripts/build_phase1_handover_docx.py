from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import build_handover_docx as base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT_ROOT / "docs" / "phase1_project_handover_manual.md"


def configure_phase1_headers(section) -> None:
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("零碳农业园区优化调度平台  |  第一阶段项目交接手册")
    base.set_run_font(run, size=8.5, color=base.MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    label = paragraph.add_run("v1.0.0  ·  ")
    base.set_run_font(label, size=8.5, color=base.MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_element = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    properties.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), base.MUTED)
    properties.append(color)
    run_element.append(properties)
    text = OxmlElement("w:t")
    text.text = "1"
    run_element.append(text)
    field.append(run_element)
    paragraph._p.append(field)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(100)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("第一阶段交付 · 版本 1.0.0")
    base.set_run_font(run, size=10, bold=True, color=base.BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("基于绿电直连的\n零碳农业园区优化调度平台")
    base.set_run_font(run, size=27, bold=True, color=base.NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(32)
    run = subtitle.add_run("项目交接手册 · 算法说明 · 逐页面演示讲稿")
    base.set_run_font(run, size=14.5, color=base.DARK_BLUE)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(76)
    run = summary.add_run("可部署  ·  可交付  ·  可维护  ·  可扩展")
    base.set_run_font(run, size=11, bold=True, color=base.BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    run = meta.add_run(f"版本日期：{date.today().isoformat()}")
    base.set_run_font(run, size=10, color=base.MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("交付分支：agent/v17-visual-analytics")
    base.set_run_font(run, size=10, color=base.MUTED)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("阅读导航", level=1)
    intro = doc.add_paragraph(
        "本手册把业务、页面、参数、算法、数据、数据库、部署、测试和现场讲稿放在同一份交接材料中。"
    )
    intro.paragraph_format.space_after = Pt(10)
    sections = [
        "第1—3章：平台目标、完成度和四个页面",
        "第4—6章：系统结构、参数来源和数据模板",
        "第7—10章：能量流、联合优化、压力测试和图表算法",
        "第11—14章：数据库、部署、测试、限制和下一阶段",
        "第15—17章：逐页面讲稿、现场问答和交接清单",
    ]
    for item in sections:
        doc.add_paragraph(item, style="List Bullet")
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    run = note.add_run("建议：")
    base.set_run_font(run, bold=True, color=base.DARK_BLUE)
    base.add_inline_runs(note, "演示人员先读第15—16章；工程师从第4章开始按数据流和验收清单接手。")
    doc.add_page_break()


def build(output: Path) -> Path:
    doc = Document()
    base.configure_document(doc)
    configure_phase1_headers(doc.sections[0])
    add_cover(doc)
    add_contents(doc)
    base.add_markdown(doc, SOURCE)

    doc.core_properties.title = "零碳农业园区优化调度平台第一阶段项目交接手册"
    doc.core_properties.subject = "业务边界、参数来源、算法实现、部署维护和逐页面演示"
    doc.core_properties.author = "项目交付组"
    doc.core_properties.keywords = "农业园区, 绿电直连, 零碳, 优化调度, 项目交接"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="生成第一阶段农业零碳项目交接手册")
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "deliverables" / "零碳农业园区优化平台_第一阶段项目交接手册.docx",
    )
    args = parser.parse_args()
    print(build(args.output.resolve()))


if __name__ == "__main__":
    main()
