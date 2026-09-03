from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_handover_docx as base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MANUAL_SOURCE = PROJECT_ROOT / "docs" / "v2_project_handover_manual.md"
SCRIPT_SOURCE = PROJECT_ROOT / "docs" / "v2_demo_script.md"
NAVIGATION_SOURCE = PROJECT_ROOT / "frontend" / "src" / "config" / "navigation.ts"
SCREENSHOT_DIR = PROJECT_ROOT / ".ui_qa_v2"

sys.path.insert(0, str(PROJECT_ROOT))
from src.core.agri_parameter_registry import PHASE1_PARAMETER_REGISTRY  # noqa: E402


GROUP_META = {
    "monitoring": ("实时监测", "OverviewPage、DomainDataPage；总览、资产与告警接口", "状态聚合、能量平衡、阈值与持续时间规则"),
    "production": ("农业生产", "DomainDataPage；农业时序与任务数据模型", "农业日历、时间窗、完成量、设备与工艺约束"),
    "forecast": ("预测中心", "AnalysisDataPage；ForecastService", "季节时段特征、趋势基线、误差评价与漂移比较"),
    "dispatch": ("计划调度", "AnalysisDataPage；能量流、策略与任务接口", "96点日前计划、15分钟滚动窗口、SOC与生产硬约束"),
    "planning": ("规划分析", "AnalysisDataPage；AgriZeroCarbonOptimizer", "容量搜索、年度时序、年化成本、可靠性与零碳约束"),
    "green-direct": ("绿电直连", "AnalysisDataPage；GreenDirectBenefitService", "送受端损耗、合同偏差、容量费、收益归因与可靠性"),
    "carbon": ("碳管理", "AnalysisDataPage；策略与参数接口", "活动数据乘排放因子、绿电属性分账、基准差额法"),
    "assets": ("资产管理", "DomainDataPage；资产登记与拓扑接口", "资产状态机、计量映射、拓扑关系与绩效指标"),
    "data": ("数据管理", "AnalysisDataPage；参数、资产与版本数据", "来源分级、质量规则、版本快照、量程与平衡校验"),
    "reports": ("报表与审计", "WorkspacePage；策略、任务与审计接口", "口径汇总、版本固化、校验和与受控导出"),
    "system": ("系统管理", "SystemStatusPage；认证、任务、监控与备份", "PBKDF2、令牌、角色权限、租约重试、健康检查与审计"),
}


SPECIAL_METHODS = {
    "M02-01": "关口与分项功率求和、最大需量、越限阈值和电能平衡残差。",
    "M02-02": "实发功率与装机、辐照和温度归一化比较，计算可用率、限发和故障损失。",
    "M02-03": "逐时计算送端注入、线路损耗、受端送达、合同偏差和中断时长。",
    "M02-04": "SOC状态递推、充放电互斥、效率、温度和等效循环校验。",
    "M02-05": "四类农业负荷分解、并发峰值、日电量和可调比例统计。",
    "M02-06": "设备状态、环境量程、变化率和工艺联锁联合判断。",
    "M02-07": "阈值加持续时间触发，按等级、对象和责任人形成闭环状态机。",
    "M03-02": "地块水量、泵站流量、扬程、时间窗和单位水量能耗约束。",
    "M03-03": "批次连续性、最短运行时间、含水率目标和单位产量能耗。",
    "M03-04": "温湿度带、热惯性、融霜窗口、库存风险和制冷移峰边界。",
    "M03-07": "基线负荷减硬约束负荷，按持续时间、恢复时间和生产影响形成柔性包络。",
    "M04-01": "季节、时段和农业任务特征预测总负荷与分项负荷，并计算误差。",
    "M04-02": "辐照驱动的物理功率基线叠加设备可用率与限发修正。",
    "M04-03": "送端资源、合同曲线、通道损耗、可用率和中断概率组合。",
    "M04-05": "分时电价、需量费用和偏差成本形成调度价格向量。",
    "M04-07": "平均绝对误差、均方根误差、平均绝对百分比误差、基线对比和漂移。",
    "M05-01": "带时间窗和设备占用的农业任务排程，人工锁定优先。",
    "M05-02": "未来24小时96点供需平衡、SOC、备用、成本和碳排联合计划。",
    "M05-03": "每15分钟滚动重算未来窗口，冻结已执行决策并保留回退方案。",
    "M05-04": "储能SOC、功率、能量、来源属性、衰减成本和备用容量联合调度。",
    "M05-05": "只移动生产允许的柔性部分，保持任务完成量和恢复约束。",
    "M05-07": "计划减实测形成偏差，按预测、设备、通信、人工和合同原因归因。",
    "M06-02": "联合搜索光伏、储能功率、储能能量和直连合同容量。",
    "M06-03": "对推荐容量执行全年时序，汇总能量、SOC、弃电、未供电和需量。",
    "M06-04": "现金流、年化成本、净现值、回收期和内部收益率适用性判断。",
    "M06-05": "同一负荷、价格和边界下比较五策略，避免跨口径比较。",
    "M06-06": "单因素和多因素扰动，计算目标和约束对参数变化的响应。",
    "M06-07": "固定容量施加七类压力场景，暴露成本、碳排和供电风险。",
    "M07-04": "合同曲线、最低消纳、免考核区间、超限偏差、容量费和可用率结算。",
    "M07-05": "实测送达与合同计划逐点比较，执行偏差考核和归因。",
    "M07-06": "物理电量、储能来源和凭证唯一标识核对，防止重复计算。",
    "M07-07": "成本变化分项归因并强制与总成本变化精确对账。",
    "M07-08": "资源不足、线路中断、合同偏差、设备故障和价格风险矩阵。",
    "M08-01": "组织边界、运营边界和计量边界版本化保存。",
    "M08-02": "普通电网净购电量乘适用排放因子，绿电属性单独核验。",
    "M08-03": "按来源、时间、计量点和凭证形成可追溯碳账本。",
    "M08-05": "基准排放减方案排放，附加性、重复性和证据完整性检查。",
    "M09-04": "缺失、越界、冻结、突变、时钟、平衡和业务一致性规则评分。",
    "M09-06": "参数编码、单位、值域、来源等级、核验状态和设置方法统一登记。",
    "M10-04": "固化输入快照、模型版本、求解状态、容量、成本、碳和风险。",
    "M10-05": "预测、计划、指令、回执、实测和参数校准形成闭环复盘。",
    "M10-06": "证据文件、日志、审批、签名、版本与SHA-256校验和关联。",
    "M11-03": "角色、菜单、操作和数据权限矩阵，API服务端执行最终校验。",
    "M11-04": "模型登记、验证、发布、灰度、监控和回退状态机。",
    "M11-08": "作业原子领取、幂等键、Worker租约、心跳、超时恢复和最大重试。",
    "M11-09": "存活、就绪、接口时延、作业计数、存储和结构化日志监控。",
    "M11-10": "SQLite在线备份、成果文件打包、SHA-256清单和维护窗口恢复。",
    "M11-11": "按主体、时间、对象、动作、结果和请求标识保存审计事件。",
}


def configure_headers(section) -> None:
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("零碳农业园区融合版V2  |  项目交接与逐页面演示手册")
    base.set_run_font(run, size=8.5, color=base.MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    label = footer.add_run("融合版V2  ·  ")
    base.set_run_font(label, size=8.5, color=base.MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_element = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    properties.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), base.MUTED)
    properties.append(color)
    run_element.append(properties)
    text = OxmlElement("w:t")
    text.text = "1"
    run_element.append(text)
    field.append(run_element)
    footer._p.append(field)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(88)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("第一阶段正式交付  ·  融合版V2")
    base.set_run_font(run, size=10, bold=True, color=base.BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)
    run = title.add_run("基于绿电直连的\n零碳农业园区优化调度平台")
    base.set_run_font(run, size=27, bold=True, color=base.NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("项目交接手册  ·  算法说明  ·  78页面对照  ·  现场演示讲稿")
    base.set_run_font(run, size=14, color=base.DARK_BLUE)

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.space_after = Pt(72)
    run = statement.add_run("可部署  ·  可交付  ·  可维护  ·  可扩展")
    base.set_run_font(run, size=11, bold=True, color=base.BLUE)

    for line in (
        f"版本日期：{date.today().isoformat()}",
        "交付分支：agent/phase2-real-data-intake",
        "数据口径：农业园区第一阶段演示基准，正式项目需按来源台账替换",
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(line)
        base.set_run_font(run, size=9.5, color=base.MUTED)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("阅读导航", level=1)
    intro = doc.add_paragraph(
        "本文件把业务边界、输入参数、算法实现、代码结构、接口、数据库、部署、测试、78个页面和现场讲稿合并为一份工程交接材料。"
    )
    intro.paragraph_format.space_after = Pt(10)
    items = [
        "第一部分：项目定位、农业与零碳边界、参数来源、算法、架构、部署和验收",
        "第二部分：平台关键页面视觉对照",
        "第三部分：15分钟主线、30分钟扩展、模块讲法、问答和故障应急",
        "附录A：78个页面的实现、算法和演示动作对照",
        "附录B：79项参数的值、单位、来源、核验状态和设置方法",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    run = note.add_run("建议顺序：")
    base.set_run_font(run, bold=True, color=base.DARK_BLUE)
    base.add_inline_runs(note, "演示人员先读第三部分；工程师按第一部分和两个附录完成交接与复核。")
    doc.add_page_break()


def add_part_title(doc: Document, text: str, *, page_break: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(20)
    run = paragraph.add_run(text)
    base.set_run_font(run, size=19, bold=True, color=base.NAVY)


def add_visual_gallery(doc: Document) -> None:
    add_part_title(doc, "第二部分  ·  融合版关键页面视觉对照", page_break=True)
    images = [
        ("overview-live.png", "图1  态势总览：能源、农业、绿电、成本、碳排和事件统一入口"),
        ("dispatch-live.png", "图2  计划调度：日前与滚动决策、储能和农业柔性负荷"),
        ("planning-step8.png", "图3  规划分析：容量、年度运行、经济性、方案和风险"),
        ("carbon-step8.png", "图4  碳管理：边界、活动数据、因子、台账和目标"),
        ("parameters-step8.png", "图5  参数中心：来源分级、核验状态和设置方法"),
        ("M11-09-step9.png", "图6  系统监控：接口、作业、存储和运维状态"),
    ]
    for index, (filename, caption_text) in enumerate(images):
        path = SCREENSHOT_DIR / filename
        if not path.exists():
            continue
        if index:
            doc.add_page_break()
        picture = doc.add_picture(str(path), width=Inches(6.25))
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture._inline.docPr.set("descr", caption_text)
        picture._inline.docPr.set("title", caption_text.split("：", 1)[0])
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(6)
        caption.paragraph_format.space_after = Pt(8)
        run = caption.add_run(caption_text)
        base.set_run_font(run, size=9.5, bold=True, color=base.DARK_BLUE)
        note = doc.add_paragraph(
            "视觉验收关注：页面标题与导航一致、人民币单位统一、关键口径有来源标签、图表与表格使用同一数据接口、无英文占位和无权限越界。"
        )
        note.paragraph_format.space_after = Pt(0)


def parse_navigation() -> list[tuple[str, str, str, str]]:
    source = NAVIGATION_SOURCE.read_text(encoding="utf-8")
    groups: list[tuple[str, str, str, str]] = []
    pattern = re.compile(
        r'createGroup\("([^"]+)",\s*"([^"]+)",\s*"[^"]+",\s*\[(.*?)\]\),',
        re.S,
    )
    page_pattern = re.compile(r'\["(M\d{2}-\d{2})",\s*"([^"]+)",\s*"([^"]+)"\]')
    for group_key, group_title, body in pattern.findall(source):
        for code, title, description in page_pattern.findall(body):
            groups.append((group_key, group_title, code, title + "\n" + description))
    return groups


def method_for(group_key: str, code: str) -> str:
    if code in SPECIAL_METHODS:
        return SPECIAL_METHODS[code]
    return GROUP_META[group_key][2] + "。"


def action_for(group_key: str, title: str) -> str:
    actions = {
        "monitoring": "先看新鲜度和告警，再看功率、电量与边界。",
        "production": "选择任务或设备，说明时间窗、完成量和不可破坏约束。",
        "forecast": "切换实际、预测和误差，说明版本、更新时间和置信边界。",
        "dispatch": "展示计划、约束、影响和回退，不直接宣称自动下发。",
        "planning": "比较容量、成本、碳排和压力风险，说明参数版本。",
        "green-direct": "从送端、通道、受端、合同、结算到证据逐层讲解。",
        "carbon": "先讲边界和因子，再讲活动数据、凭证和核查。",
        "assets": "定位资产、计量点、状态和责任人，强调待现场核验。",
        "data": "查看来源等级、质量规则、版本和影响范围。",
        "reports": "说明口径、版本、审计证据和受控导出。",
        "system": "只展示状态和权限，不暴露口令、密钥或完整个人信息。",
    }
    return actions[group_key]


def add_page_appendix(doc: Document) -> None:
    add_part_title(doc, "附录A  ·  78个页面实现与算法对照", page_break=True)
    intro = doc.add_paragraph(
        "态势总览为M01；其余77页来自导航配置。表中“实现”说明页面与服务职责，“算法或规则”说明核心计算，“演示动作”给出演示时的操作重点。"
    )
    intro.paragraph_format.space_after = Pt(10)

    base.add_table(
        doc,
        [
            ["页面", "实现与数据", "算法或规则", "演示动作"],
            [
                "M01\n态势总览",
                "OverviewPage读取总览、策略和系统摘要；卡片、能量流、趋势、农业分项与事件联动。",
                "年度时序聚合、逐时能量平衡、策略基准和数据新鲜度判断。",
                "从顶部指标到能量流，再到农业分项；强调同一模型口径和人民币单位。",
            ],
        ],
    )

    grouped_pages: dict[str, tuple[str, list[tuple[str, str]]]] = {}
    for group_key, group_title, code, title_and_description in parse_navigation():
        grouped_pages.setdefault(group_key, (group_title, []))[1].append((code, title_and_description))

    for group_key, (group_title, pages) in grouped_pages.items():
        heading = doc.add_heading(group_title, level=2)
        heading.paragraph_format.page_break_before = True
        rows = [["页面", "实现与数据", "算法或规则", "演示动作"]]
        for code, title_and_description in pages:
            title, description = title_and_description.split("\n", 1)
            implementation = GROUP_META[group_key][1] + "。" + description
            rows.append(
                [
                    f"{code}\n{title}",
                    implementation,
                    method_for(group_key, code),
                    action_for(group_key, title),
                ]
            )
        base.add_table(doc, rows)


def format_parameter_value(parameter) -> str:
    if parameter.value is None:
        return "待补充"
    value = float(parameter.value)
    if abs(value) >= 1000:
        return f"{value:,.2f}"
    if value.is_integer():
        return f"{value:.0f}"
    return f"{value:.4f}".rstrip("0").rstrip(".")


def add_parameter_appendix(doc: Document) -> None:
    add_part_title(doc, "附录B  ·  79项参数来源台账", page_break=True)
    counts = Counter(parameter.source_grade.value for parameter in PHASE1_PARAMETER_REGISTRY)
    summary = "；".join(f"{grade}：{count}项" for grade, count in sorted(counts.items()))
    paragraph = doc.add_paragraph(
        f"台账共{len(PHASE1_PARAMETER_REGISTRY)}项。来源分布为{summary}。E级值仅用于演示，正式项目必须按来源引用和设置方法替换并重新核验。"
    )
    paragraph.paragraph_format.space_after = Pt(10)

    groups: dict[str, list] = {}
    for parameter in PHASE1_PARAMETER_REGISTRY:
        groups.setdefault(parameter.group, []).append(parameter)

    for group, parameters in groups.items():
        doc.add_heading(group, level=2)
        rows = [["编码与参数", "当前值与单位", "来源与状态", "来源引用及设置方法"]]
        for parameter in parameters:
            range_text = ""
            if parameter.minimum is not None or parameter.maximum is not None:
                lower = "无" if parameter.minimum is None else format(float(parameter.minimum), "g")
                upper = "无" if parameter.maximum is None else format(float(parameter.maximum), "g")
                range_text = f"\n范围：{lower}—{upper}"
            rows.append(
                [
                    f"{parameter.code}\n{parameter.name}",
                    f"{format_parameter_value(parameter)} {parameter.unit}{range_text}",
                    f"{parameter.source_grade.value}\n{parameter.verification_status.value}",
                    f"{parameter.source_reference}\n设置：{parameter.setting_method}",
                ]
            )
        base.add_table(doc, rows)


def build(output: Path) -> Path:
    doc = Document()
    base.configure_document(doc)
    configure_headers(doc.sections[0])
    add_cover(doc)
    add_contents(doc)

    add_part_title(doc, "第一部分  ·  项目交接手册")
    base.add_markdown(doc, MANUAL_SOURCE)
    add_visual_gallery(doc)

    add_part_title(doc, "第三部分  ·  逐页面演示讲稿", page_break=True)
    base.add_markdown(doc, SCRIPT_SOURCE)
    add_page_appendix(doc)
    add_parameter_appendix(doc)

    doc.core_properties.title = "基于绿电直连的零碳农业园区优化调度平台融合版V2项目交接与演示手册"
    doc.core_properties.subject = "业务边界、参数来源、算法实现、78页面、部署维护和现场演示"
    doc.core_properties.author = "项目交付组"
    doc.core_properties.keywords = "零碳农业园区, 绿电直连, V17, 优化调度, 项目交接, 演示讲稿"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="生成融合版V2项目交接与演示Word手册")
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "deliverables" / "零碳农业园区融合版V2_项目交接与演示手册.docx",
    )
    args = parser.parse_args()
    print(build(args.output.resolve()))


if __name__ == "__main__":
    main()
